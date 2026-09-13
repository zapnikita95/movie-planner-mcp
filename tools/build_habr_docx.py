from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ARTICLE = ROOT / "HABR_ARTICLE.md"
COVER = ROOT / "assets" / "movie-planner-mcp-cover.png"
OUT_DIR = ROOT / "dist"
OUT = OUT_DIR / "Movie Planner MCP - статья для Хабра.docx"


def set_run_font(run, name: str = "Aptos", size: int | None = None, bold: bool | None = None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_paragraph_spacing(paragraph, before: int = 0, after: int = 8, line: float = 1.08):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_inline_markdown(paragraph, text: str, *, size: int = 11):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, "Consolas", size)
            run.font.color.rgb = RGBColor(35, 35, 35)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size)


def add_code_block(doc: Document, code: str):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F3F4F6")
    p = cell.paragraphs[0]
    p.paragraph_format.left_indent = Cm(0.15)
    p.paragraph_format.right_indent = Cm(0.15)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(code.strip())
    set_run_font(run, "Consolas", 10)
    run.font.color.rgb = RGBColor(20, 20, 20)


def add_body_paragraph(doc: Document, text: str):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=9, line=1.12)
    add_inline_markdown(p, text, size=11)


def build_docx():
    OUT_DIR.mkdir(exist_ok=True)
    source = ARTICLE.read_text(encoding="utf-8").splitlines()
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    styles["Normal"].font.size = Pt(11)
    for style_name, size in [("Title", 24), ("Heading 1", 18), ("Heading 2", 14)]:
        st = styles[style_name]
        st.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
        st._element.rPr.rFonts.set(qn("w:ascii"), st.font.name)
        st._element.rPr.rFonts.set(qn("w:hAnsi"), st.font.name)
        st._element.rPr.rFonts.set(qn("w:eastAsia"), st.font.name)
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0, 0, 0)

    title_done = False
    in_code = False
    code_lines: list[str] = []

    for raw in source:
        line = raw.rstrip()
        if line.startswith("!["):
            if COVER.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(COVER), width=Cm(16.5))
                set_paragraph_spacing(p, after=14)
            continue
        if line.strip() == "```text":
            in_code = True
            code_lines = []
            continue
        if line.strip() == "```" and in_code:
            add_code_block(doc, "\n".join(code_lines))
            in_code = False
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not line.strip():
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(line[2:].strip())
            set_run_font(run, "Aptos Display", 24, True)
            set_paragraph_spacing(p, after=10)
            title_done = True
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1" if title_done else "Heading 2")
            run = p.add_run(line[3:].strip())
            set_run_font(run, "Aptos", 15, True)
            set_paragraph_spacing(p, before=12, after=5)
            continue
        add_body_paragraph(doc, line)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Movie Planner MCP · черновик статьи для Хабра")
    set_run_font(run, "Aptos", 9)
    run.font.color.rgb = RGBColor(90, 90, 90)

    doc.core_properties.title = "Movie Planner MCP статья для Хабра"
    doc.core_properties.subject = "Черновик статьи о Movie Planner MCP"
    doc.core_properties.author = "Movie Planner"
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_docx())
